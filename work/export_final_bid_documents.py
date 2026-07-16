from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "assets" / "js" / "bid-document-flow-data.js"
OUTPUT_DIR = ROOT / "assets" / "bid-documents"

# narrative_proposal is the closest document archetype. The user's global Word
# specification is the authority override for A4 geometry, black headings and
# Normal paragraph rhythm. Source-page content is otherwise preserved verbatim.
CONTENT_WIDTH_DXA = 8311
TABLE_INDENT_DXA = 120
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_TOP_BOTTOM_CM = 2.54
MARGIN_LEFT_RIGHT_CM = 3.17
HEADER_DISTANCE_CM = 1.50
FOOTER_DISTANCE_CM = 1.75


def load_data() -> dict:
    text = DATA_FILE.read_text(encoding="utf-8-sig")
    prefix = "window.BidDocumentFlowData="
    payload = text[text.index(prefix) + len(prefix):].strip().rstrip(";")
    return json.loads(payload)


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def set_run_font(run, size: float = 12, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "宋体"
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_outline_level(style, level: int) -> None:
    ppr = style.element.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), str(level - 1))


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading_tokens = {
        1: (22, 17, 16.5, 2.4, 21.6),
        2: (16, 13, 13, 1.7208, 28.75),
        3: (16, 13, 13, 1.7208, 36),
        4: (14, 14, 14.5, 1.55, 43.2),
        5: (14, 14, 14.5, 1.55, 50.4),
        6: (12, 12, 3.2, 1.3208, 57.55),
        7: (12, 12, 3.2, 1.3208, 64.8),
        8: (12, 12, 3.2, 1.3208, 72),
        9: (12, 12, 3.2, 1.3208, 79.15),
    }
    for level, (size, before, after, line_spacing, indent) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        fmt = style.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = line_spacing
        fmt.left_indent = Pt(indent)
        fmt.first_line_indent = Pt(-indent)
        fmt.keep_with_next = True
        fmt.keep_together = True
        set_outline_level(style, level)

    caption = styles["Caption"]
    caption.font.name = "宋体"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10.5)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.line_spacing = 1.0

    list_style = styles["List Bullet"]
    list_style.font.name = "宋体"
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    list_style.font.size = Pt(12)
    list_style.paragraph_format.left_indent = Cm(0.75)
    list_style.paragraph_format.first_line_indent = Cm(-0.35)
    list_style.paragraph_format.line_spacing = 1.5
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)


def configure_section(section) -> None:
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(MARGIN_TOP_BOTTOM_CM)
    section.bottom_margin = Cm(MARGIN_TOP_BOTTOM_CM)
    section.left_margin = Cm(MARGIN_LEFT_RIGHT_CM)
    section.right_margin = Cm(MARGIN_LEFT_RIGHT_CM)
    section.header_distance = Cm(HEADER_DISTANCE_CM)
    section.footer_distance = Cm(FOOTER_DISTANCE_CM)


def picture_stream(source: Path, prefer_jpeg: bool = False) -> tuple[BytesIO, int, int]:
    stream = BytesIO()
    with Image.open(source) as image:
        width, height = image.size
        if prefer_jpeg:
            image.convert("RGB").save(stream, format="JPEG", quality=93, subsampling=0, optimize=True)
        else:
            has_alpha = image.mode in {"RGBA", "LA"} or "transparency" in image.info
            image.convert("RGBA" if has_alpha else "RGB").save(stream, format="PNG", optimize=True)
    stream.seek(0)
    return stream, width, height


def clear_header(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    for paragraph in header.paragraphs:
        for run in list(paragraph.runs):
            paragraph._p.remove(run._r)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)


def add_stamp_header(section, source: Path) -> None:
    clear_header(section)
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.line_spacing = 1.0
    stream, _, _ = picture_stream(source)
    shape = paragraph.add_run().add_picture(stream, width=Cm(1.9))
    shape._inline.docPr.set("descr", "企业公章")
    shape._inline.docPr.set("title", "企业公章")


def resolve_asset(src: str) -> Path:
    source = (ROOT / "pages" / Path(src.replace("/", "\\"))).resolve()
    if not source.is_relative_to(ROOT):
        raise ValueError(f"asset outside workspace: {src}")
    if not source.exists():
        raise FileNotFoundError(source)
    return source


def append_inline(paragraph, node, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    if node.text:
        run = paragraph.add_run(node.text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    for child in node:
        tag = str(child.tag).lower() if isinstance(child.tag, str) else ""
        if tag == "br":
            paragraph.add_run().add_break()
        else:
            append_inline(
                paragraph,
                child,
                size=size,
                bold=bold or tag in {"b", "strong"},
                italic=italic or tag in {"i", "em"},
            )
        if child.tail:
            run = paragraph.add_run(child.tail)
            set_run_font(run, size=size, bold=bold, italic=italic)


def add_paragraph_from_html(doc: Document, node) -> None:
    classes = set((node.get("class") or "").split())
    paragraph = doc.add_paragraph(style="Normal")
    if "toc-entry" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0 if "level-1" in classes else 0.65 if "level-2" in classes else 1.3)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(14.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        title_nodes = node.xpath("./span")
        page_nodes = node.xpath("./b")
        if title_nodes and page_nodes:
            append_inline(paragraph, title_nodes[0], size=11.5, bold="level-1" in classes)
            paragraph.add_run().add_tab()
            append_inline(paragraph, page_nodes[0], size=11.5, bold=True)
            return
    if "doc-signature-line" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
    append_inline(paragraph, node, size=12)


def add_heading_from_html(doc: Document, node, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    append_inline(paragraph, node, size={1: 22, 2: 16, 3: 16, 4: 14}[level], bold=True)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table_from_html(doc: Document, node) -> None:
    source_rows = node.xpath("./thead/tr | ./tbody/tr | ./tr")
    column_count = max(len(row.xpath("./th | ./td")) for row in source_rows)
    base_width = CONTENT_WIDTH_DXA // column_count
    widths = [base_width] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    table = doc.add_table(rows=len(source_rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, source_row in enumerate(source_rows):
        source_cells = source_row.xpath("./th | ./td")
        is_header = any(str(cell.tag).lower() == "th" for cell in source_cells)
        if is_header:
            tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        for column_index, source_cell in enumerate(source_cells):
            cell = table.rows[row_index].cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if is_header:
                shade_cell(cell, "F2F2F2")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header or column_count >= 4 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            append_inline(paragraph, source_cell, size=9.5 if column_count >= 6 else 10.5, bold=is_header)
    set_table_geometry(table, widths)


def add_image_from_html(doc: Document, node, keep_with_caption: bool = False) -> None:
    source = resolve_asset(node.get("src"))
    is_full_scan = source.parent.name == "business" and source.stem.startswith("page-")
    stream, pixel_width, pixel_height = picture_stream(source, prefer_jpeg=is_full_scan)
    max_width_cm = 13.8 if is_full_scan else 14.0
    max_height_cm = 21.0 if is_full_scan else 20.5
    aspect = pixel_width / pixel_height
    width_cm = max_width_cm
    height_cm = width_cm / aspect
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * aspect
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = keep_with_caption
    shape = paragraph.add_run().add_picture(stream, width=Cm(width_cm), height=Cm(height_cm))
    alt_text = clean_text(node.get("alt")) or source.stem
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)


def add_figure_from_html(doc: Document, node) -> None:
    caption_nodes = node.xpath("./figcaption")
    images = node.xpath("./img")
    for index, image in enumerate(images):
        add_image_from_html(doc, image, keep_with_caption=bool(caption_nodes) and index == len(images) - 1)
    for caption_node in caption_nodes:
        paragraph = doc.add_paragraph(style="Caption")
        append_inline(paragraph, caption_node, size=10.5)


def add_list_from_html(doc: Document, node) -> None:
    for item in node.xpath("./li"):
        paragraph = doc.add_paragraph(style="List Bullet")
        append_inline(paragraph, item, size=12)


def add_page_content(doc: Document, page_html: str, stats: dict[str, int]) -> None:
    root = html.fragment_fromstring(page_html, create_parent="div")
    for node in root:
        tag = str(node.tag).lower() if isinstance(node.tag, str) else ""
        if tag == "img" and "doc-stamp" in (node.get("class") or "").split():
            stats["stamps"] += 1
        elif tag in {"h1", "h2", "h3", "h4"}:
            stats["headings"] += 1
            add_heading_from_html(doc, node, int(tag[1]))
        elif tag == "p":
            stats["paragraphs"] += 1
            add_paragraph_from_html(doc, node)
        elif tag == "table":
            stats["tables"] += 1
            add_table_from_html(doc, node)
        elif tag == "figure":
            stats["figures"] += len(node.xpath("./img"))
            add_figure_from_html(doc, node)
        elif tag == "ul":
            stats["list_items"] += len(node.xpath("./li"))
            add_list_from_html(doc, node)


def build_document(kind: str, data: dict) -> tuple[Path, dict[str, int]]:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.core_properties.title = data["name"]
    doc.core_properties.subject = ""
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""

    stamp = OUTPUT_DIR / kind / "figures" / "company-stamp.webp"
    add_stamp_header(doc.sections[0], stamp)
    stats = {"pages": len(data["pages"]), "paragraphs": 0, "headings": 0, "tables": 0, "figures": 0, "list_items": 0, "stamps": 0}

    for index, page in enumerate(data["pages"], start=1):
        add_page_content(doc, page["html"], stats)
        if kind == "business" and index == 7:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            configure_section(section)
            clear_header(section)
        elif index < len(data["pages"]):
            doc.add_page_break()

    path = OUTPUT_DIR / f"{kind}.docx"
    doc.save(path)
    return path, stats


def main() -> None:
    source = load_data()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for kind in ("technical", "business"):
        path, stats = build_document(kind, source[kind])
        print(f"{path.name}\t{path.stat().st_size}\t{json.dumps(stats, ensure_ascii=False)}")


if __name__ == "__main__":
    main()
