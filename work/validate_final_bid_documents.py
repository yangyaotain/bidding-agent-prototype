from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree, html


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "assets" / "js" / "bid-document-flow-data.js"
OUTPUT_DIR = ROOT / "assets" / "bid-documents"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
FORBIDDEN = ("可编辑 Word 示例文件", "Word 示例", "静态原型", "投标助手", "本文件用于")


def load_data() -> dict:
    text = DATA_FILE.read_text(encoding="utf-8-sig")
    prefix = "window.BidDocumentFlowData="
    return json.loads(text[text.index(prefix) + len(prefix):].strip().rstrip(";"))


def normalized(value: str) -> str:
    return re.sub(r"\s+", "", value)


def digest(value: str) -> str:
    return hashlib.sha256(normalized(value).encode("utf-8")).hexdigest()


def source_body_text(data: dict) -> str:
    pieces: list[str] = []
    for page in data["pages"]:
        root = html.fragment_fromstring(page["html"], create_parent="div")
        for node in root:
            tag = str(node.tag).lower() if isinstance(node.tag, str) else ""
            if tag in {"h1", "h2", "h3", "h4", "p", "table", "ul"}:
                pieces.append(node.text_content())
            elif tag == "figure":
                pieces.extend(caption.text_content() for caption in node.xpath("./figcaption"))
    return "".join(pieces)


def source_counts(data: dict) -> dict[str, int]:
    counts = {"pages": len(data["pages"]), "headings": 0, "tables": 0, "figures": 0, "stamps": 0}
    for page in data["pages"]:
        root = html.fragment_fromstring(page["html"], create_parent="div")
        counts["headings"] += len(root.xpath("./h1 | ./h2 | ./h3 | ./h4"))
        counts["tables"] += len(root.xpath("./table"))
        counts["figures"] += len(root.xpath("./figure/img"))
        counts["stamps"] += len(root.xpath('./img[@class="doc-stamp"]'))
    return counts


def validate_table_geometry(doc: Document) -> None:
    for index, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa"
        assert int(tbl_w.get(qn("w:w"))) == 8311
        assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == 120
        grid = [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid.findall(qn("w:gridCol"))]
        assert sum(grid) == 8311
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert tr_pr.find(qn("w:cantSplit")) is not None
            assert tr_pr.find(qn("w:trHeight")) is None
            widths = [int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w"))) for cell in row.cells]
            assert widths == grid, f"table {index}: row widths differ from grid"


def validate(kind: str, data: dict) -> None:
    path = OUTPUT_DIR / f"{kind}.docx"
    counts = source_counts(data)
    with ZipFile(path) as package:
        assert package.testzip() is None
        names = set(package.namelist())
        assert {"word/document.xml", "word/styles.xml", "docProps/core.xml"}.issubset(names)
        document_root = etree.fromstring(package.read("word/document.xml"))
        generated_text = "".join(document_root.xpath(".//w:body//w:t/text()", namespaces=NS))
        assert digest(generated_text) == digest(source_body_text(data)), f"{kind}: source/export text mismatch"
        page_breaks = len(document_root.xpath('.//w:br[@w:type="page"]', namespaces=NS))
        expected_page_breaks = counts["pages"] - 1 if kind == "technical" else counts["pages"] - 2
        assert page_breaks == expected_page_breaks
        assert not document_root.xpath(".//wp:anchor", namespaces=NS)
        document_doc_pr = document_root.xpath(".//wp:docPr", namespaces=NS)
        assert len(document_doc_pr) == counts["figures"]
        assert all(node.get("descr") and node.get("title") for node in document_doc_pr)

        header_names = sorted(name for name in names if name.startswith("word/header") and name.endswith(".xml"))
        assert header_names
        header_images = 0
        for header_name in header_names:
            header_root = etree.fromstring(package.read(header_name))
            assert not header_root.xpath(".//w:t[text()]", namespaces=NS)
            assert not header_root.xpath(".//wp:anchor", namespaces=NS)
            header_doc_pr = header_root.xpath(".//wp:docPr", namespaces=NS)
            header_images += len(header_doc_pr)
            assert all(node.get("descr") and node.get("title") for node in header_doc_pr)
        assert header_images == 1

        for footer_name in (name for name in names if name.startswith("word/footer") and name.endswith(".xml")):
            footer_root = etree.fromstring(package.read(footer_name))
            assert not footer_root.xpath(".//w:t[text()]", namespaces=NS)

        all_xml_text = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.endswith(".xml")
        )
        for forbidden in FORBIDDEN:
            assert forbidden not in all_xml_text

    doc = Document(path)
    expected_sections = 1 if kind == "technical" else 2
    assert len(doc.sections) == expected_sections
    for section in doc.sections:
        assert abs(section.page_width.cm - 21.0) < 0.03
        assert abs(section.page_height.cm - 29.7) < 0.03
        assert abs(section.top_margin.cm - 2.54) < 0.03
        assert abs(section.bottom_margin.cm - 2.54) < 0.03
        assert abs(section.left_margin.cm - 3.17) < 0.03
        assert abs(section.right_margin.cm - 3.17) < 0.03
        assert abs(section.header_distance.cm - 1.50) < 0.03
        assert abs(section.footer_distance.cm - 1.75) < 0.03

    normal_style = doc.styles["Normal"]
    assert normal_style.font.size == Pt(12)
    assert normal_style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert normal_style.paragraph_format.first_line_indent == Pt(24)
    assert normal_style.paragraph_format.line_spacing == 1.5
    heading_count = sum(1 for paragraph in doc.paragraphs if paragraph.style.name.startswith("Heading "))
    assert heading_count == counts["headings"]
    assert len(doc.tables) == counts["tables"]
    assert len(doc.inline_shapes) == counts["figures"]
    validate_table_geometry(doc)

    print(
        f"PASS {path.name}: source_pages={counts['pages']}, headings={heading_count}, "
        f"tables={len(doc.tables)}, inline_images={len(doc.inline_shapes)}, bytes={path.stat().st_size}"
    )


if __name__ == "__main__":
    source = load_data()
    validate("technical", source["technical"])
    validate("business", source["business"])
